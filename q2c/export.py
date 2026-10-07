"""
Quote-to-close workbook in the PM's layout (Summary / Pivot / Quote to close / Raw data), every number a
live formula:
  Raw data        - the quote lines used (values)
  Sales lines     - the invoice lines of those Fact #s (values)
  Quote to close  - one row per Fact #: Count, Quoted Qty, AVG, Sales Qty, Status, Invoiced value are FORMULAS
  Pivot / By Area - SUMIFS over Quote to close
  Summary         - the PM's Summary rows as formulas
Changing the sales window on the Parameters sheet recalculates everything.
Pure functions - no Streamlit calls.
"""
import io
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BLUE = PatternFill("solid", fgColor="0D68A3")
WHITE_BOLD = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14, color="0D68A3")
EDIT = PatternFill("solid", fgColor="FFF8E1")
MT, PCT, DATE, INT = "#,##0.00", "0.0%", "DD MMM YYYY", "#,##0"
QTC = "'Quote to close'"
RAW = "'Raw data'"
SL = "'Sales lines'"


def _crit(text: str) -> str:
    """Excel criteria literal: escape wildcards (* ? ~) and quotes so a name only matches itself."""
    t = str(text).replace("~", "~~").replace("*", "~*").replace("?", "~?").replace('"', '""')
    return f'"{t}"'


def _header(ws, row, headers, widths=None):
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=j, value=h)
        c.fill, c.font = BLUE, WHITE_BOLD
        c.alignment = Alignment(wrap_text=True, vertical="center")
    if widths:
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w


def _pivot_sheet(wb, name, title, key_col, keys):
    ws = wb.create_sheet(name)
    ws["A1"], ws["A1"].font = title, TITLE
    _header(ws, 3, [name.replace("By ", "") if name != "Pivot" else "Customer", "Sum of Count",
                    "Sum of Quoted Qty", "Sum of AVG MT / Quote", "Sum of Sales Qty", "Fill rate (Sales / Quoted)",
                    "Sum of Invoiced value"], [55, 12, 16, 18, 16, 16, 18])
    col = f"{QTC}!${key_col}:${key_col}"
    r = 4
    for k in keys:
        cr = _crit(k)
        ws.cell(row=r, column=1, value=str(k))
        ws.cell(row=r, column=2, value=f"=SUMIFS({QTC}!$D:$D,{col},{cr})")
        ws.cell(row=r, column=3, value=f"=SUMIFS({QTC}!$E:$E,{col},{cr})")
        ws.cell(row=r, column=4, value=f"=SUMIFS({QTC}!$F:$F,{col},{cr})")
        ws.cell(row=r, column=5, value=f"=SUMIFS({QTC}!$G:$G,{col},{cr})")
        ws.cell(row=r, column=6, value=f'=IFERROR(E{r}/C{r},"")')
        ws.cell(row=r, column=7, value=f"=SUMIFS({QTC}!$K:$K,{col},{cr})")
        r += 1
    ws.cell(row=r, column=1, value="Grand Total")
    for j, L in zip(range(2, 6), "BCDE"):
        ws.cell(row=r, column=j, value=f"=SUM({L}4:{L}{r - 1})")
    ws.cell(row=r, column=6, value=f'=IFERROR(E{r}/C{r},"")')
    ws.cell(row=r, column=7, value=f"=SUM(G4:G{r - 1})")
    for rr in range(4, r + 1):
        ws.cell(row=rr, column=2).number_format = INT
        for j in (3, 4, 5):
            ws.cell(row=rr, column=j).number_format = MT
        ws.cell(row=rr, column=6).number_format = PCT
        ws.cell(row=rr, column=7).number_format = "#,##0"
    for j in range(1, 8):
        ws.cell(row=r, column=j).font = BOLD
    ws.freeze_panes = "A4"


def build_workbook(table: pd.DataFrame, quote_lines: pd.DataFrame, sales_lines: pd.DataFrame,
                   params, pivot_customers: list, pivot_areas: list, files_note: str) -> bytes:
    wb = Workbook()

    # ── Summary (the PM's rows) ──
    ws = wb.active
    ws.title = "Summary"
    ws["A1"], ws["A1"].font = "Quote-to-Close Summary", TITLE
    ws["A2"] = (f"Quotes {params.quote_start:%d %b %Y} – {params.quote_end:%d %b %Y} · invoices "
                f"{params.sales_start:%d %b %Y} – {params.sales_end:%d %b %Y}  |  {files_note}")
    _header(ws, 4, ["Key Metric", "Value", "% of total"], [48, 16, 12])
    rows = [
        (5, "Total Quotations (line items)", f"=COUNTA({QTC}!$C:$C)-1", INT, None),
        (6, "Quotations Closed (>= 1 order)", f'=COUNTIF({QTC}!$J:$J,"Closed")', INT, None),
        (7, "Quotations Lost (zero sales)", f'=COUNTIF({QTC}!$J:$J,"Lost")', INT, None),
        (8, "Quote-to-Close Ratio (by count)", "=IFERROR(B6/B5,0)", PCT, None),
        (10, "Total Quoted Volume", f"=SUM({QTC}!$E:$E)", MT, None),
        (12, "Quotation closed (by volume)", f'=SUMIFS({QTC}!$E:$E,{QTC}!$J:$J,"Closed")', MT, "=IFERROR(B12/B10,0)"),
        (13, "Quotation lost (by volume)", f'=SUMIFS({QTC}!$E:$E,{QTC}!$J:$J,"Lost")', MT, "=IFERROR(B13/B10,0)"),
        (14, "Total Sales Volume Realised", f"=SUM({QTC}!$G:$G)", MT, None),
        (15, "Volume Fill Rate (Sales / Quoted volume)", "=IFERROR(B14/B10,0)", PCT, None),
        (16, "Volume Fill Rate (Sales / Quotation closed)", "=IFERROR(B14/B12,0)", PCT, None),
        (18, "Sum of AVG Quoted volume", f"=SUM({QTC}!$F:$F)", MT, None),
        (19, "Quotation closed (by AVG volume)", f'=SUMIFS({QTC}!$F:$F,{QTC}!$J:$J,"Closed")', MT, "=IFERROR(B19/B18,0)"),
        (20, "Quotation lost (by AVG volume)", f'=SUMIFS({QTC}!$F:$F,{QTC}!$J:$J,"Lost")', MT, "=IFERROR(B20/B18,0)"),
        (22, "Total Invoiced Value (Invoice Amount)", f"=SUM({QTC}!$K:$K)", "#,##0", None),
    ]
    for r, label, f, fmt, pct in rows:
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=f).number_format = fmt
        if pct:
            ws.cell(row=r, column=3, value=pct).number_format = PCT
    ws["A24"] = "Sales window complete?"
    ws["B24"] = "=Parameters!B8"

    # ── How calculated ──
    hc = wb.create_sheet("How calculated")
    hc["A1"], hc["A1"].font = "How every number is calculated (the PM's method)", TITLE
    _header(hc, 3, ["Number", "Definition"], [40, 110])
    defs = [
        ("Quotation (line item)", "One Fact # quoted in the quote period. All order types and salesmen are included."),
        ("Area", "The customer's Area in the quote export, except the customers re-labelled by the PM "
                 "(Petrochemical, Consumer Area, Central Procurement, …)."),
        ("Customer", "The customer on the Fact #'s first quote line."),
        ("Count", "Number of quote lines for the Fact # = COUNTIFS(Raw data Fact #)."),
        ("Quoted Qty", "SUM of MT of all its quote lines = SUMIFS(Raw data MT)."),
        ("AVG MT / Quote", "Quoted Qty ÷ Count."),
        ("Sales Qty", "SUM of invoiced MT (Shipped Tons) for the Fact # with Invoice date inside the sales window "
                      "(Parameters B4–B5) = SUMIFS(Sales lines MT). 'Fact Tic NB' in the invoice files is the Fact #."),
        ("Invoiced value", "SUM of Invoice Amount of the same invoice lines."),
        ("Closed / Lost", "Closed if Sales Qty > 0, Lost if Sales Qty = 0."),
        ("Quote-to-Close Ratio (by count)", "Closed ÷ Total quotations."),
        ("Quotation closed / lost (by volume)", "Quoted Qty of the closed / lost Fact #s, and its % of Total Quoted Volume."),
        ("Volume Fill Rate (Sales / Quoted volume)", "Total Sales ÷ Total Quoted Volume."),
        ("Volume Fill Rate (Sales / Quotation closed)", "Total Sales ÷ Quoted Qty of the closed Fact #s."),
        ("Quotation closed / lost (by AVG volume)", "The same split using AVG MT / Quote instead of Quoted Qty."),
        ("Last quote # / Date", "The Trans # and Issue Date of the Fact #'s latest quote line."),
    ]
    for r, (a, b) in enumerate(defs, start=4):
        hc.cell(row=r, column=1, value=a).font = BOLD
        hc.cell(row=r, column=2, value=b).alignment = Alignment(wrap_text=True, vertical="top")

    # ── Parameters ──
    wp = wb.create_sheet("Parameters")
    prm = [("Parameter", "Value", "Meaning"),
           ("Quote period start", params.quote_start.to_pydatetime(), "Raw data sheet already filtered to this period"),
           ("Quote period end", params.quote_end.to_pydatetime(), ""),
           ("Sales window start (invoice date)", params.sales_start.to_pydatetime(), "Editable"),
           ("Sales window end (invoice date)", params.sales_end.to_pydatetime(), "Editable"),
           ("Last invoice date in the files", params.export_day.to_pydatetime(), "")]
    for i, r in enumerate(prm, start=1):
        for j, v in enumerate(r, start=1):
            wp.cell(row=i, column=j, value=v)
    _header(wp, 1, prm[0], [34, 16, 60])
    for r in range(2, 7):
        wp.cell(row=r, column=2).number_format = DATE
    for r in (4, 5):
        wp.cell(row=r, column=2).fill = EDIT
    wp["A8"] = "Sales window complete?"
    wp["B8"] = '=IF(B5<=B6,"Yes","No - window ends after the last invoice date")'

    # ── Quote to close (formulas) ──
    wq = wb.create_sheet("Quote to close")
    _header(wq, 1, ["Area", "Customer", "Fact #", "Count", "Quoted Qty", "AVG MT / Quote", "Sales Qty",
                    "Last quote #", "Date", "Status", "Invoiced value"], [18, 45, 12, 8, 12, 14, 12, 13, 13, 9, 14])
    for r, rec in enumerate(table.itertuples(index=False), start=2):
        wq.append([rec.area, rec.customer, rec.item,
                   f"=COUNTIFS({RAW}!$B:$B,$C{r})",
                   f"=SUMIFS({RAW}!$E:$E,{RAW}!$B:$B,$C{r})",
                   f"=IFERROR(E{r}/D{r},0)",
                   f'=SUMIFS({SL}!$D:$D,{SL}!$A:$A,$C{r},{SL}!$C:$C,">="&Parameters!$B$4,{SL}!$C:$C,"<="&Parameters!$B$5)',
                   rec.last_quote_no, rec.last_quote.to_pydatetime(),
                   f'=IF(G{r}>0,"Closed","Lost")',
                   f'=SUMIFS({SL}!$E:$E,{SL}!$A:$A,$C{r},{SL}!$C:$C,">="&Parameters!$B$4,{SL}!$C:$C,"<="&Parameters!$B$5)'])
        for j, fmt in ((5, MT), (6, MT), (7, MT), (9, DATE), (11, "#,##0")):
            wq.cell(row=r, column=j).number_format = fmt
    wq.freeze_panes = "A2"
    wq.auto_filter.ref = f"A1:K{len(table) + 1}"

    # ── Pivots ──
    _pivot_sheet(wb, "Pivot", "Quote-to-Close by Customer (fill rate, highest first)", "B", pivot_customers)
    _pivot_sheet(wb, "By Area", "Quote-to-Close by Area", "A", pivot_areas)

    # ── Raw data ──
    wr = wb.create_sheet("Raw data")
    _header(wr, 1, ["Trans #", "Fact #", "Issue Date", "Customer", "MT", "Area", "Order Type", "Salesman"],
            [12, 12, 13, 45, 10, 18, 16, 24])
    for rec in quote_lines.itertuples(index=False):
        wr.append([rec.quote_no, rec.item, rec.quote_day.to_pydatetime(), rec.customer, float(rec.mt), rec.area,
                   rec.order_type, rec.salesman])
    for r in range(2, len(quote_lines) + 2):
        wr.cell(row=r, column=3).number_format = DATE
        wr.cell(row=r, column=5).number_format = MT
    wr.freeze_panes = "A2"
    wr.auto_filter.ref = f"A1:H{len(quote_lines) + 1}"

    # ── Sales lines (invoices) ──
    ws2 = wb.create_sheet("Sales lines")
    _header(ws2, 1, ["Fact #", "Invoice #", "Invoice date", "MT (Shipped Tons)", "Invoice Amount", "Customer",
                     "MFG Code", "SC #"], [12, 15, 13, 14, 14, 45, 12, 12])
    for rec in sales_lines.itertuples(index=False):
        ws2.append([rec.item, rec.inv_no, rec.sale_day.to_pydatetime(), float(rec.mt), float(rec.amount),
                    rec.customer, rec.mfg_code, rec.sc_no])
    for r in range(2, len(sales_lines) + 2):
        ws2.cell(row=r, column=3).number_format = DATE
        ws2.cell(row=r, column=4).number_format = "#,##0.000"
        ws2.cell(row=r, column=5).number_format = "#,##0.00"
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:H{len(sales_lines) + 1}"

    order = ["Summary", "Pivot", "By Area", "Quote to close", "Raw data", "Sales lines", "Parameters", "How calculated"]
    wb._sheets = [wb[n] for n in order]
    wb.active = 0
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
