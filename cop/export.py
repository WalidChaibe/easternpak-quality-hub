"""Download-to-Excel of every COP strategy table. Pure functions - no Streamlit calls."""
import io
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BLUE = PatternFill("solid", fgColor="0D68A3")
WHITE_BOLD = Font(bold=True, color="FFFFFF")


def customer_columns(cy: int, py: int, years: list) -> dict:
    """Internal column -> header, in display order."""
    cols = {"customer": "Customer", "status": "Status", "strategy": "Strategy", "grow_tag": "COP growth driver",
            "action": "What to do"}
    for y in years:
        cols[f"mt_{y}"] = f"MT {y}"
    cols["vol_pct"] = "Vol Δ%"
    for y in years:
        cols[f"price_{y}"] = f"Price $/t {y}"
    cols.update({"dprice": "Price Δ $/t", f"cost_{py}": f"Cost $/t {py}", f"cost_{cy}": f"Cost $/t {cy}",
                 "dcost": "Cost Δ $/t"})
    for y in years:
        cols[f"cop_{y}"] = f"COP {y} (SAR k)"
    for y in years:
        cols[f"copt_{y}"] = f"COP $/t {y}"
    cols.update({"dcopt": "COP $/t Δ", "dcop": "ΔCOP (SAR k)", "capture": "Price change kept as COP %", "reading": "Price → COP reading",
                 "vol_eff": "Volume effect", "price_eff": "Price effect", "cost_eff": "Cost effect",
                 "vol_move": "Volume move", "price_move": "Price/t move", "cost_move": "Cost/t move",
                 "copt_move": "COP/t move",
                 "margin_tier": "Margin tier", "key_account": "Key account", "key_reason": "Key account because",
                 "worth": "Worth fighting for", "trend": "Multi-year COP trend",
                 "recent_copt_cy": f"COP $/t last months {cy}", "recent_copt_py": f"COP $/t same months {py}",
                 "momentum": "Momentum", "share_cy_cop": f"Share of {cy} COP", "at_stake": "COP at stake (SAR k)"})
    return cols


def build_workbook(sheets: dict) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        for name, df in sheets.items():
            df.to_excel(w, sheet_name=name[:31], index=False)
            ws = w.sheets[name[:31]]
            for c in ws[1]:
                c.fill, c.font = BLUE, WHITE_BOLD
                c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.freeze_panes = "B2"
            for j, col in enumerate(df.columns, start=1):
                width = 45 if str(col) in ("Customer", "What to do", "Meaning", "Who falls here", "Reading") else 14
                ws.column_dimensions[get_column_letter(j)].width = width
                if pd.api.types.is_float_dtype(df[col]):
                    fmt = "0.0%" if ("%" in str(col) or "Share" in str(col)) else "#,##0.00"
                    for r in range(2, len(df) + 2):
                        ws.cell(row=r, column=j).number_format = fmt
    return buf.getvalue()
