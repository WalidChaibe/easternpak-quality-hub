"""
Forecast workbook with LIVE formulas, so the PM can adjust customers in Excel.

  Customers sheet  - forecast MT per customer per month (plain numbers: the PM edits these)
                     + Total (formula) + Model total (fixed copy) + Change vs model (formula)
  FT sheet         - every FT cell = FT share x that customer's month on the Customers sheet
                     (looked up by customer name, so sorting/filtering Customers is safe)
  Monthly total    - SUM of each Customers month column
Pure functions - no Streamlit calls.
"""
import io
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="0D68A3")
HEADER_FONT = Font(bold=True, color="FFFFFF")
EDIT_FILL = PatternFill("solid", fgColor="FFF8E1")     # cells the PM is expected to edit
MT_FMT = "#,##0.0"
PCT_FMT = "0.0%"


def _header(ws, headers):
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2"


def _widths(ws, widths: dict):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def build_workbook(customers: pd.DataFrame, fc: pd.DataFrame, mix: pd.DataFrame,
                   fc_months: pd.PeriodIndex, cross_check: pd.Series, eff_days: pd.Series,
                   partial_month: pd.Period | None, notes: list[str]) -> bytes:
    """
    customers: index = customer; columns industry, area, model, history_months, size_change.
    fc: customer x fc_months forecast MT (partial month already includes MT invoiced so far).
    mix: output of models.ft_mix (customer, ft, product_name, share, months_invoiced, last_invoiced).
    """
    month_labels = [m.strftime("%b %Y") + ("*" if m == partial_month else "") for m in fc_months]
    n = len(fc_months)
    wb = Workbook()

    # ── Customers ──
    ws = wb.active
    ws.title = "Customers"
    meta_cols = ["Customer", "Industry", "Area", "Model", "History months", "Size change"]
    m0 = len(meta_cols) + 1                       # first month column (G)
    _header(ws, meta_cols + month_labels + ["Total MT", "Model total MT", "Change vs model MT"])
    order = fc.sum(axis=1).sort_values(ascending=False).index
    for r, cust in enumerate(order, start=2):
        info = customers.loc[cust] if cust in customers.index else None
        row = [cust,
               info["industry"] if info is not None else "",
               info["area"] if info is not None else "",
               info["model"] if info is not None else "",
               info["history_months"] if info is not None else None,
               "Yes" if (info is not None and bool(info["size_change"])) else ""]
        vals = [float(v) for v in fc.loc[cust, fc_months].values]
        first, last = get_column_letter(m0), get_column_letter(m0 + n - 1)
        tot_col = get_column_letter(m0 + n)
        model_col = get_column_letter(m0 + n + 1)
        ws.append(row + vals + [f"=SUM({first}{r}:{last}{r})", sum(vals), f"={tot_col}{r}-{model_col}{r}"])
        for j in range(n + 3):
            cell = ws.cell(row=r, column=m0 + j)
            cell.number_format = MT_FMT
            if j < n:
                cell.fill = EDIT_FILL
    last_row = len(order) + 1
    ws.auto_filter.ref = f"A1:{get_column_letter(m0 + n + 2)}{last_row}"
    _widths(ws, {"A": 42, "B": 22, "C": 14, "D": 30, "E": 9, "F": 8})
    for j in range(n + 3):
        ws.column_dimensions[get_column_letter(m0 + j)].width = 11

    # ── FT ──
    wf = wb.create_sheet("FT")
    ft_meta = ["Customer", "FT", "Product name", "Model", "Share", "Months invoiced (last 12)",
               "Last invoiced", "Customer row"]
    f0 = len(ft_meta) + 1                         # first month column (I)
    _header(wf, ft_meta + month_labels + ["Total MT"])
    for r, rec in enumerate(mix.itertuples(index=False), start=2):
        model = customers.loc[rec.customer, "model"] if rec.customer in customers.index else ""
        row = [rec.customer, rec.ft, rec.product_name, model, float(rec.share), int(rec.months_invoiced),
               rec.last_invoiced.to_pydatetime(), f"=MATCH($A{r},Customers!$A:$A,0)"]
        formulas = [f"=$E{r}*INDEX(Customers!{get_column_letter(m0 + j)}:{get_column_letter(m0 + j)},$H{r})"
                    for j in range(n)]
        wf.append(row + formulas + [f"=SUM({get_column_letter(f0)}{r}:{get_column_letter(f0 + n - 1)}{r})"])
        wf.cell(row=r, column=5).number_format = PCT_FMT
        wf.cell(row=r, column=7).number_format = "DD MMM YYYY"
        for j in range(n + 1):
            wf.cell(row=r, column=f0 + j).number_format = MT_FMT
    wf.auto_filter.ref = f"A1:{get_column_letter(f0 + n)}{len(mix) + 1}"
    wf.column_dimensions["H"].hidden = True       # helper column for the lookup
    _widths(wf, {"A": 42, "B": 10, "C": 40, "D": 30, "E": 8, "F": 11, "G": 13})
    for j in range(n + 1):
        wf.column_dimensions[get_column_letter(f0 + j)].width = 11

    # ── Monthly total ──
    wm = wb.create_sheet("Monthly total")
    _header(wm, ["Month", "Forecast MT (live, from Customers)", "Model forecast MT (fixed)",
                 "Cross-check: company-level model MT", "Effective invoicing days"])
    for i, m in enumerate(fc_months):
        col = get_column_letter(m0 + i)
        r = i + 2
        wm.append([month_labels[i], f"=SUM(Customers!{col}:{col})",
                   float(fc[m].sum()), float(cross_check[m]), float(eff_days[m])])
        for c in (2, 3, 4):
            wm.cell(row=r, column=c).number_format = MT_FMT
        wm.cell(row=r, column=5).number_format = "0.0"
    tr = n + 2
    wm.append(["Total", f"=SUM(B2:B{tr - 1})", f"=SUM(C2:C{tr - 1})", f"=SUM(D2:D{tr - 1})", ""])
    for c in range(1, 5):
        wm.cell(row=tr, column=c).font = Font(bold=True)
        if c > 1:
            wm.cell(row=tr, column=c).number_format = MT_FMT
    _widths(wm, {"A": 12, "B": 22, "C": 20, "D": 22, "E": 14})

    # ── Notes ──
    wn = wb.create_sheet("Notes")
    for line in notes:
        wn.append([line])
    wn.column_dimensions["A"].width = 130

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
